import hashlib
import html
import json
import math
import sys
from pathlib import Path

import pymupdf as fitz


def emit(value):
    print(json.dumps(value, ensure_ascii=False), flush=True)


def progress(page, total):
    emit({'progress': page, 'total': total})


def source_hash(path):
    with open(path, 'rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def page_at(doc, number):
    if type(number) is not int or not 1 <= number <= len(doc):
        raise ValueError('Invalid page number.')
    return doc[number - 1]


def assert_permission(doc, modify=False):
    if doc.needs_pass:
        raise ValueError('Encrypted PDF requires a password; this engine does not accept passwords.')
    permission = fitz.PDF_PERM_MODIFY if modify else fitz.PDF_PERM_COPY
    if not doc.permissions & permission:
        raise ValueError('Document permission does not allow this operation.')
    if modify:
        if doc.get_sigflags() > 0:
            raise ValueError('Signed PDF cannot be rewritten by this engine.')
        for page in doc:
            if any(widget.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE for widget in page.widgets()):
                raise ValueError('PDF contains signature fields; rewriting is disabled.')
            if any(annot.type[0] == fitz.PDF_ANNOT_REDACT for annot in (page.annots() or [])):
                raise ValueError('PDF has pending redactions. Apply or remove them before editing.')
        for xref in range(1, doc.xref_length()):
            if doc.xref_get_key(xref, 'ByteRange')[0] != 'null':
                raise ValueError('PDF contains a signature byte range; rewriting is disabled.')


def blocks_for(page):
    blocks = []
    invisible = [fitz.Rect(span['bbox']) for span in page.get_texttrace() if span['type'] == 3]
    for block in page.get_text('dict', flags=fitz.TEXTFLAGS_TEXT, sort=True)['blocks']:
        if block['type'] != 0:
            continue
        lines = block['lines']
        spans = [span for line in lines for span in line['spans']]
        if not spans:
            continue
        horizontal = all(abs(line['dir'][0] - 1) < .001 and abs(line['dir'][1]) < .001 for line in lines)
        blocks.append({
            'id': block['number'],
            'rect': list(block['bbox']),
            'text': '\n'.join(''.join(span['text'] for span in line['spans']) for line in lines),
            'fontSize': round(sum(span['size'] for span in spans) / len(spans), 2),
            'color': '#%06x' % spans[0]['color'],
            'editable': horizontal and not any(fitz.Rect(block['bbox']).intersects(rect) for rect in invisible),
            'spans': [list(span['bbox']) for span in spans],
        })
    return blocks


def inspect(doc, request):
    page = page_at(doc, request['page'])
    blocks = blocks_for(page)
    for block in blocks:
        block.pop('spans')
    return {'blocks': blocks, 'width': page.cropbox.width, 'height': page.cropbox.height}


def edit(doc, request):
    assert_permission(doc, modify=True)
    page = page_at(doc, request['page'])
    blocks = blocks_for(page)
    block = next((item for item in blocks if item['id'] == request.get('blockId')), None)
    if not block or not block['editable']:
        raise ValueError('Selected text block is unavailable or has unsupported text direction.')
    text = request.get('text')
    size = request.get('fontSize')
    if not isinstance(text, str) or len(text) > 20000:
        raise ValueError('Text must contain at most 20000 characters.')
    if isinstance(size, bool) or not isinstance(size, (int, float)) or not math.isfinite(size) or not 4 <= size <= 96:
        raise ValueError('Font size must be 4–96 pt.')
    rect = fitz.Rect(block['rect'])
    for other in blocks:
        if other['id'] != block['id'] and any(fitz.Rect(span).intersects(fitz.Rect(target)) for span in other['spans'] for target in block['spans']):
            raise ValueError('Text overlaps another block; independent editing is unsafe.')
    markup = '<div>' + html.escape(text).replace('\n', '<br>') + '</div>'
    css = 'div { margin:0; padding:0; font-family:sans-serif; font-size:%spt; color:%s; }' % (size, block['color'])
    # Preflight with the same layout engine so overflow never deletes the original text.
    if text.strip():
        with fitz.open() as preview:
            test = preview.new_page(width=page.cropbox.width, height=page.cropbox.height)
            spare, _ = test.insert_htmlbox(rect, markup, css=css, scale_low=1)
            if spare < 0:
                raise ValueError('Text overflows this block. Shorten it or reduce font size; cross-block reflow is not supported.')
    for span in block['spans']:
        page.add_redact_annot(fitz.Rect(span), fill=False, cross_out=False)
    page.apply_redactions(images=0, graphics=0, text=0)
    if text.strip():
        spare, _ = page.insert_htmlbox(rect, markup, css=css, scale_low=1)
        if spare < 0:
            raise ValueError('Text layout failed; no output has been saved.')
    doc.subset_fonts()
    doc.save(request['output'], garbage=4, deflate=True)
    return {'message': 'Original text removed and block reflowed; font family may be substituted.'}


def ocr(doc, request):
    assert_permission(doc, modify=True)
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    engine = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)
    recognized = 0
    processed = 0
    for index, page in enumerate(doc):
        if page.get_text().strip():
            progress(index + 1, len(doc))
            continue
        # Inference uses the displayed page; map all four corners back to unrotated page points.
        scale = min(2.5, math.sqrt(12_000_000 / (page.rect.width * page.rect.height)))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False, annots=False)
        image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)[:, :, ::-1].copy()
        results, _ = engine(image)
        del image, pix
        count = 0
        font = fitz.Font('china-s')
        for points, text, confidence in results or []:
            if not text.strip() or confidence < .5:
                continue
            mapped = [fitz.Point(float(x) / scale, float(y) / scale) * page.derotation_matrix for x, y in points]
            rect = fitz.Rect(min(p.x for p in mapped), min(p.y for p in mapped), max(p.x for p in mapped), max(p.y for p in mapped))
            if rect.width <= 0 or rect.height <= 0:
                continue
            size = min(rect.height * .8, rect.width / max(font.text_length(text, fontsize=1), .1))
            if size <= 0:
                continue
            page.insert_text((rect.x0, rect.y0 + rect.height * .8), text, fontsize=size, fontname='china-s', render_mode=3)
            count += 1
        recognized += count
        processed += bool(count)
        progress(index + 1, len(doc))
    if not recognized:
        raise ValueError('No new text recognized. Existing text pages are skipped; low-confidence or unsupported scans may need another OCR model.')
    doc.subset_fonts()
    doc.save(request['output'], garbage=3, deflate=True)
    return {'message': 'OCR: %s lines on %s pages. Search layer added; scanned image content is unchanged.' % (recognized, processed)}


def convert(doc, request):
    assert_permission(doc)
    kind = request['kind']
    output = request['output']
    if kind == 'docx':
        if not any(page.get_text().strip() for page in doc):
            raise ValueError('Scanned document has no text layer. Run OCR first.')
        from pdf2docx import Converter
        converter = Converter(request['input'])
        try:
            converter.convert(output, multi_processing=False)
        finally:
            converter.close()
        return {'message': 'DOCX reconstructed with layout, images and tables. Review font substitution and complex layouts.'}
    if kind == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        book.remove(book.active)
        total = 0
        for index, page in enumerate(doc):
            for table_index, table in enumerate(page.find_tables().tables):
                sheet = book.create_sheet('P%s-T%s' % (index + 1, table_index + 1))
                from openpyxl.styles import Font, Alignment
                for row_number, row in enumerate(table.extract(), 1):
                    # Treat PDF content as literal data, never executable spreadsheet formulas.
                    for column, value in enumerate(row, 1):
                        cell = sheet.cell(row_number, column)
                        cell.value = value
                        cell.font = Font(name='Arial', size=11)
                        cell.alignment = Alignment(wrap_text=True, vertical='top')
                        if isinstance(value, str):
                            cell.data_type = 's'
                total += 1
            progress(index + 1, len(doc))
        if not total:
            raise ValueError('No detectable tables. XLSX export is for table extraction, not full page reconstruction.')
        book.save(output)
        return {'message': '%s tables exported; verify merged cells and borderless tables.' % total}
    if kind == 'pptx':
        from io import BytesIO
        from pptx import Presentation
        from pptx.util import Pt
        presentation = Presentation()
        presentation.slide_width = Pt(doc[0].rect.width)
        presentation.slide_height = Pt(doc[0].rect.height)
        for index, page in enumerate(doc):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            scale = min(2, math.sqrt(8_000_000 / (page.rect.width * page.rect.height)))
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            ratio = min(presentation.slide_width / page.rect.width, presentation.slide_height / page.rect.height)
            width, height = int(page.rect.width * ratio), int(page.rect.height * ratio)
            slide.shapes.add_picture(BytesIO(pix.tobytes('png')), int((presentation.slide_width - width) / 2), int((presentation.slide_height - height) / 2), width, height)
            del pix
            progress(index + 1, len(doc))
        presentation.save(output)
        return {'message': 'PPTX uses page images to preserve appearance. Text and tables are not editable slide objects.'}
    raise ValueError('Unsupported conversion kind.')


def main():
    request = json.loads(sys.stdin.readline().lstrip('\ufeff'))
    if request.get('action') == 'capabilities':
        import importlib.util
        return {'version': fitz.__version__, 'ocr': bool(importlib.util.find_spec('rapidocr_onnxruntime')), 'docx': bool(importlib.util.find_spec('pdf2docx')), 'xlsx': bool(importlib.util.find_spec('openpyxl')), 'pptx': bool(importlib.util.find_spec('pptx'))}
    path = Path(request['input'])
    if path.stat().st_size > 512 * 1024 * 1024:
        raise ValueError('Advanced processing is limited to 512 MB input.')
    digest = source_hash(path)
    if request.get('sourceHash') and request['sourceHash'] != digest:
        raise ValueError('PDF changed since block inspection. Read blocks again.')
    with fitz.open(path) as doc:
        if not doc.is_pdf or len(doc) == 0 or len(doc) > 2000:
            raise ValueError('Expected a PDF with 1–2000 pages.')
        action = request['action']
        if action == 'inspect':
            assert_permission(doc)
            return {**inspect(doc, request), 'sourceHash': digest}
        if action == 'edit':
            return edit(doc, request)
        if action == 'ocr':
            return ocr(doc, request)
        if action == 'convert':
            return convert(doc, request)
        raise ValueError('Unsupported engine action.')


if __name__ == '__main__':
    try:
        emit({'result': main()})
    except Exception as error:
        emit({'error': str(error)})
        sys.exit(1)
