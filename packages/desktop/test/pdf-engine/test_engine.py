import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ENGINE = Path(__file__).resolve().parents[2] / 'static/pdf-engine/engine.py'


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.workspace = tempfile.TemporaryDirectory(prefix='yarktext-test-')
        self.root = Path(self.workspace.name)
        self.input = self.root / 'input.pdf'
        with fitz.open() as doc:
            page = doc.new_page(width=595, height=842)
            page.insert_textbox((40, 40, 420, 130), 'Original paragraph with enough width.\nSecond line for reflow testing.', fontsize=12)
            page.insert_text((40, 210), 'Unrelated text stays intact.', fontsize=12)
            for x in [40, 180, 320]:
                page.draw_line((x, 300), (x, 390))
            for y in [300, 330, 360, 390]:
                page.draw_line((40, y), (320, y))
            for row, values in enumerate([['Name', 'Value'], ['Alpha', '123'], ['=1+1', '456']]):
                for column, value in enumerate(values):
                    page.insert_text((50 + column * 140, 320 + row * 30), value, fontsize=10)
            image = Image.new('RGB', (60, 40), '#7c3aed')
            data = io.BytesIO()
            image.save(data, format='PNG')
            page.insert_image((400, 300, 520, 380), stream=data.getvalue())
            doc.save(self.input)

    def tearDown(self):
        self.workspace.cleanup()

    def invoke(self, action, **options):
        request = {'action': action, 'input': str(self.input), **options}
        process = subprocess.run([sys.executable, str(ENGINE)], input=json.dumps(request), capture_output=True, text=True, encoding='utf-8', timeout=180)
        entries = []
        for line in process.stdout.splitlines():
            try:
                entries.append(json.loads(line))
            except ValueError:
                pass
        final = entries[-1] if entries else {'error': process.stderr}
        return process.returncode, final

    def test_edit_removes_original_and_preserves_other_objects(self):
        _, inspected = self.invoke('inspect', page=1)
        block = inspected['result']['blocks'][0]
        output = self.root / 'edited.pdf'
        code, result = self.invoke('edit', page=1, blockId=block['id'], sourceHash=inspected['result']['sourceHash'], text='Rewritten 中文 paragraph.', fontSize=10, output=str(output))
        self.assertEqual(code, 0, result)
        with fitz.open(output) as doc:
            text = doc[0].get_text()
            self.assertIn('Rewritten', text)
            self.assertIn('中文', text)
            self.assertNotIn('Original paragraph', text)
            self.assertIn('Unrelated text', text)
            self.assertEqual(len(doc[0].get_images()), 1)

    def test_overflow_and_stale_inspection_rejected(self):
        _, inspected = self.invoke('inspect', page=1)
        block = inspected['result']['blocks'][0]
        output = self.root / 'invalid.pdf'
        code, _ = self.invoke('edit', page=1, blockId=block['id'], text='overflow ' * 2000, fontSize=96, sourceHash=inspected['result']['sourceHash'], output=str(output))
        self.assertNotEqual(code, 0)
        self.assertFalse(output.exists())
        code, _ = self.invoke('edit', page=1, blockId=block['id'], text='Changed', fontSize=10, sourceHash='0' * 64, output=str(output))
        self.assertNotEqual(code, 0)

    def test_docx_reconstructs_tables_images_and_text(self):
        output = self.root / 'out.docx'
        code, result = self.invoke('convert', kind='docx', output=str(output))
        self.assertEqual(code, 0, result)
        doc = Document(output)
        self.assertIn('Original paragraph', '\n'.join(p.text for p in doc.paragraphs))
        self.assertGreaterEqual(len(doc.tables), 1)
        self.assertGreaterEqual(len(doc.inline_shapes), 1)

    def test_xlsx_is_literal_data_not_formulas(self):
        output = self.root / 'out.xlsx'
        code, result = self.invoke('convert', kind='xlsx', output=str(output))
        self.assertEqual(code, 0, result)
        book = load_workbook(output)
        self.assertEqual(book.worksheets[0]['A1'].value, 'Name')
        cells = [cell for sheet in book for row in sheet for cell in row]
        self.assertTrue(any(cell.value == '=1+1' for cell in cells))
        self.assertFalse(any(cell.data_type == 'f' for cell in cells))
        book.close()

    def test_pptx_keeps_page_dimensions_and_picture(self):
        output = self.root / 'out.pptx'
        code, result = self.invoke('convert', kind='pptx', output=str(output))
        self.assertEqual(code, 0, result)
        deck = Presentation(output)
        self.assertEqual(len(deck.slides), 1)
        self.assertEqual(len(deck.slides[0].shapes), 1)
        self.assertAlmostEqual(deck.slide_width / 12700, 595)

    def test_scan_ocr_adds_searchable_text_without_changing_pixels(self):
        image = Image.new('RGB', (1200, 800), 'white')
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 48)
        draw.text((80, 120), 'YarkText OCR 12345', fill='black', font=font)
        data = io.BytesIO()
        image.save(data, format='PNG')
        scan = self.root / 'scan.pdf'
        with fitz.open() as doc:
            page = doc.new_page(width=600, height=400)
            page.insert_image(page.rect, stream=data.getvalue())
            doc.save(scan)
        self.input = scan
        output = self.root / 'ocr.pdf'
        code, result = self.invoke('ocr', output=str(output))
        self.assertEqual(code, 0, result)
        with fitz.open(scan) as before, fitz.open(output) as after:
            self.assertIn('12345', after[0].get_text())
            self.assertEqual(before[0].get_pixmap().samples, after[0].get_pixmap().samples)
        self.input = output
        code, inspected = self.invoke('inspect', page=1)
        self.assertEqual(code, 0, inspected)
        self.assertTrue(inspected['result']['blocks'])
        self.assertFalse(any(block['editable'] for block in inspected['result']['blocks']))


if __name__ == '__main__':
    unittest.main()
